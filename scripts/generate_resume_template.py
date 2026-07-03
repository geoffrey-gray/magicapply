"""Generate the canonical resume template shipped at configs/resume_template.docx.

Run once to (re)produce the template. The output is checked in so users get a
working default without needing python-docx installed for a straight
`uv run magicapply` invocation.

Usage:
    uv run python scripts/generate_resume_template.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_OUT = Path(__file__).resolve().parents[1] / "configs" / "resume_template.docx"


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def _add_text(doc: Document, text: str, *, bold: bool = False, italic: bool = False,
              align: int | None = None, size: int | None = None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def build() -> Path:
    doc = Document()

    # Identity block
    _add_text(doc, "{{ name }}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=20)
    _add_text(
        doc,
        "{{ email }} · {{ phone }} · {{ location }}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    # Summary
    _add_heading(doc, "Summary")
    _add_text(doc, "{{ summary }}")

    # Experience — loop with paragraph-level {%p %} tags so the tag paragraphs
    # disappear from the rendered output.
    _add_heading(doc, "Experience")
    _add_text(doc, "{%p for exp in experience %}")
    _add_text(doc, "{{ exp.title }} — {{ exp.company }}", bold=True)
    _add_text(doc, "{{ exp.start }} to {{ exp.end }}", italic=True)
    _add_text(doc, "{%p for bullet in exp.bullets %}")
    _add_bullet(doc, "{{ bullet }}")
    _add_text(doc, "{%p endfor %}")
    _add_text(doc, "{%p endfor %}")

    # Education
    _add_heading(doc, "Education")
    _add_text(doc, "{%p for edu in education %}")
    _add_text(doc, "{{ edu.school }} — {{ edu.degree }} {{ edu.field }}", bold=True)
    _add_text(doc, "{{ edu.graduated }}", italic=True)
    _add_text(doc, "{%p endfor %}")

    # Skills
    _add_heading(doc, "Skills")
    _add_text(doc, "{{ skills | join(', ') }}")

    _OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(_OUT))
    return _OUT


if __name__ == "__main__":
    out = build()
    print(f"wrote {out}")
