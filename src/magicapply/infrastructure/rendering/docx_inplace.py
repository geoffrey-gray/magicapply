"""Format-preserving DOCX tailoring via in-place keyword swaps.

Opens the operator's original DOCX and walks paragraphs → runs, swapping
each bank synonym for the canonical bank term wherever the JD emphasises
that term. Because we only mutate ``run.text``, every other run property
(bold / italic / font / color / size / spacing / margins / hyperlinks / …)
survives byte-for-byte. This is the Phase 1 tailoring path per
``final_dod_plan.md`` W.1.

The Phase K ``DocxResumeRenderer`` template-based renderer is retired
alongside this class — the Phase 1 operator ships their own polished
DOCX; the tool never regenerates from a generic template.
"""

from __future__ import annotations

import logging
import re
import shutil
from collections.abc import Iterable
from pathlib import Path

from docx import Document

from magicapply.config.models import KeywordEntry

logger = logging.getLogger(__name__)


class InPlaceDocxTailorer:
    """Apply bank-synonym → JD-term swaps to a DOCX in place.

    Instances are stateless; the tailorer takes all inputs on each
    ``render`` call. Case-insensitive whole-word matching. When no bank
    entry contributes a swap, the output is a byte-identical copy of the
    input.
    """

    def render(
        self,
        *,
        source_docx: Path,
        matched_bank: Iterable[KeywordEntry],
        jd_terms: Iterable[str],
        out_path: Path,
    ) -> Path:
        source_docx = Path(source_docx)
        out_path = Path(out_path)
        if not source_docx.exists():
            raise FileNotFoundError(
                f"source DOCX missing: {source_docx}. Set "
                f"BaseResume.source_docx_path to a valid file."
            )
        out_path.parent.mkdir(parents=True, exist_ok=True)

        swaps = _build_swaps(matched_bank, jd_terms)
        if not swaps:
            # No applicable swaps → byte-copy so the caller can rely on
            # `out_path` being a valid DOCX identical to the source.
            shutil.copy(source_docx, out_path)
            return out_path

        doc = Document(str(source_docx))

        # Every paragraph in the body, and every paragraph in every table
        # cell. Headers / footers are not touched in Phase 1 — the
        # operator's polished resume shouldn't have keyword-bearing text
        # there anyway; revisit if a real resume proves otherwise.
        for paragraph in doc.paragraphs:
            _apply_swaps_to_runs(paragraph.runs, swaps)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _apply_swaps_to_runs(paragraph.runs, swaps)

        doc.save(str(out_path))
        return out_path


def _build_swaps(
    matched_bank: Iterable[KeywordEntry],
    jd_terms: Iterable[str],
) -> list[tuple[re.Pattern[str], str]]:
    """Build (pattern, replacement) pairs for every applicable swap.

    An entry contributes swaps when its ``term`` appears (case-insensitive)
    in ``jd_terms``: for each synonym, we swap the synonym → term in-place.
    Longer synonyms sort first so ``multi-agent orchestration`` fires
    before a shorter overlapping synonym would; whole-word matching
    prevents ``sql`` from eating ``postgresql``.
    """
    jd_lower = {t.strip().lower() for t in jd_terms if t and t.strip()}
    if not jd_lower:
        return []

    swaps: list[tuple[re.Pattern[str], str]] = []
    for entry in matched_bank:
        term = entry.term.strip()
        if not term or term.lower() not in jd_lower:
            continue
        for synonym in sorted(entry.synonyms, key=len, reverse=True):
            syn = synonym.strip()
            if not syn or syn.lower() == term.lower():
                continue
            pattern = re.compile(
                r"\b" + re.escape(syn) + r"\b",
                re.IGNORECASE,
            )
            swaps.append((pattern, term))
            logger.debug(
                "in-place swap: %r → %r (bank entry %r)",
                syn,
                term,
                entry.term,
            )
    return swaps


def _apply_swaps_to_runs(
    runs: Iterable,
    swaps: list[tuple[re.Pattern[str], str]],
) -> None:
    """Apply every swap to every run in ``runs`` at the run.text level."""
    for run in runs:
        new_text = run.text
        for pattern, replacement in swaps:
            new_text = pattern.sub(replacement, new_text)
        if new_text != run.text:
            run.text = new_text
