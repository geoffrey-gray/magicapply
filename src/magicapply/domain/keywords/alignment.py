"""ATS-style keyword coverage: JD-extracted terms vs resume text.

Correct scoring model (operator intent):

1. Extract concrete skill/tool terms **from the job description**.
2. Count how many of those terms appear on the resume.
3. Score = ``100 * matched / len(jd_terms)`` (e.g. 4 of 10 → 40).

The KeywordBank is **not** the scoring vocabulary. It is optional here only
to credit resume synonyms (JD says ``k8s``, resume says ``kubernetes``)
when a bank entry links the forms. Tailoring still owns bank-driven DOCX
swaps separately.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from magicapply.config.models import KeywordBank, KeywordEntry
from magicapply.domain.models.job import Job
from magicapply.domain.models.resume import BaseResume, TailoredResume


@dataclass(frozen=True, slots=True)
class AlignmentResult:
    """Structured keyword-alignment outcome (score 0–100 + diagnostic lists)."""

    value: int
    rationale: str
    jd_terms: list[str] = field(default_factory=list)
    matched: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)


def phrase_in_text(phrase: str, haystack: str) -> bool:
    """Case-insensitive whole-phrase match with word boundaries.

    Multi-word phrases use ``\\b`` on each edge so ``sql`` does not match
    inside ``postgresql``, while ``distributed systems`` matches that pair.
    """
    phrase = phrase.strip()
    if not phrase:
        return False
    pattern = re.compile(
        r"\b" + re.escape(phrase.strip()) + r"\b",
        re.IGNORECASE,
    )
    return pattern.search(haystack) is not None


def jd_form_for_entry(entry: KeywordEntry, jd_text: str) -> str | None:
    """Return the longest term/synonym form present in the JD, or None."""
    forms = [entry.term, *entry.synonyms]
    present = [f.strip() for f in forms if f and phrase_in_text(f, jd_text)]
    if not present:
        return None
    # Prefer longer phrases so "deep reinforcement learning" wins over "rl".
    return max(present, key=lambda s: (len(s), s.lower()))


def normalize_jd_terms(terms: list[str]) -> list[str]:
    """Deduplicate extracted JD terms (case-insensitive, preserve first form)."""
    out: list[str] = []
    seen: set[str] = set()
    for raw in terms:
        term = " ".join(str(raw).split()).strip()
        if not term:
            continue
        key = term.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(term)
    return out


# Common JD↔resume acronym pairs (not a skill bank — just surface-form aliases).
_ACRONYM_ALIASES: dict[str, tuple[str, ...]] = {
    "machine learning": ("ml", "machine learning"),
    "ml": ("ml", "machine learning"),
    "deep learning": ("dl", "deep learning"),
    "artificial intelligence": ("ai", "artificial intelligence"),
    "ai": ("ai", "artificial intelligence"),
    "natural language processing": ("nlp", "natural language processing"),
    "nlp": ("nlp", "natural language processing"),
    "large language models": ("llm", "llms", "large language models"),
    "large language model": ("llm", "llms", "large language model"),
    "llm": ("llm", "llms", "large language model", "large language models"),
    "llms": ("llm", "llms", "large language model", "large language models"),
    "kubernetes": ("k8s", "kubernetes"),
    "k8s": ("k8s", "kubernetes"),
    "amazon web services": ("aws", "amazon web services"),
    "aws": ("aws", "amazon web services"),
    "google cloud platform": ("gcp", "google cloud", "google cloud platform"),
    "gcp": ("gcp", "google cloud", "google cloud platform"),
    "a/b testing": ("a/b testing", "ab testing", "a-b testing"),
    "ab testing": ("a/b testing", "ab testing", "a-b testing"),
}


def term_on_resume(
    term: str,
    resume_text: str,
    bank: KeywordBank | None = None,
) -> bool:
    """True if ``term`` (or a known alias / bank synonym) appears on the resume."""
    candidates = [term]
    aliases = _ACRONYM_ALIASES.get(term.lower().strip())
    if aliases:
        candidates.extend(aliases)
    for cand in candidates:
        if phrase_in_text(cand, resume_text):
            return True
    if bank is None or not bank.keywords:
        return False
    # Credit bank synonym forms: JD extracted "k8s", resume has "kubernetes".
    term_l = term.lower()
    for entry in bank.keywords:
        forms = [entry.term, *entry.synonyms]
        forms_l = [f.lower() for f in forms if f]
        if term_l not in forms_l and not any(
            term_l in f or f in term_l for f in forms_l
        ):
            continue
        if any(phrase_in_text(f, resume_text) for f in forms if f):
            return True
    return False


def score_jd_keyword_coverage(
    resume_text: str,
    jd_terms: list[str],
    *,
    bank: KeywordBank | None = None,
) -> AlignmentResult:
    """Fraction of JD-extracted keywords present on the resume.

    Example: JD yields 10 terms, resume has 4 → score 40.
    """
    terms = normalize_jd_terms(jd_terms)
    if not terms:
        return AlignmentResult(
            value=0,
            rationale="keyword alignment: no keywords extracted from JD",
            jd_terms=[],
            matched=[],
            missing=[],
        )

    matched: list[str] = []
    missing: list[str] = []
    for term in terms:
        if term_on_resume(term, resume_text, bank):
            matched.append(term)
        else:
            missing.append(term)

    value = int(round(100 * len(matched) / len(terms)))
    value = max(0, min(100, value))

    matched_s = ", ".join(matched) if matched else "—"
    missing_s = ", ".join(missing) if missing else "—"
    rationale = (
        f"keyword alignment: {len(matched)}/{len(terms)} JD keywords on resume "
        f"(matched: {matched_s}; missing: {missing_s})"
    )
    return AlignmentResult(
        value=value,
        rationale=rationale,
        jd_terms=terms,
        matched=matched,
        missing=missing,
    )


def score_keyword_alignment(
    resume_text: str,
    job: Job,
    bank: KeywordBank | None = None,
    *,
    jd_terms: list[str] | None = None,
) -> AlignmentResult:
    """Score JD→resume keyword coverage.

    Prefer explicit ``jd_terms`` from :class:`KeywordExtractor`. When omitted,
    falls back to bank terms present in the JD (legacy / unit-test helper only
    — production scoring always passes extracted terms).
    """
    if jd_terms is not None:
        return score_jd_keyword_coverage(resume_text, jd_terms, bank=bank)

    # Legacy path: bank ∩ JD (not the preferred scoring model).
    if bank is None or not bank.keywords:
        return AlignmentResult(
            value=0,
            rationale="keyword alignment: no JD terms provided and empty bank",
        )

    jd_text = "\n".join(
        part
        for part in (job.title, job.company, job.location or "", job.description)
        if part
    )
    derived: list[str] = []
    for entry in bank.keywords:
        form = jd_form_for_entry(entry, jd_text)
        if form is not None:
            derived.append(form)
    return score_jd_keyword_coverage(resume_text, derived, bank=bank)


def serialize_resume_text(resume: BaseResume | TailoredResume) -> str:
    """Plain-text blob used for keyword matching (deterministic order)."""
    lines = [f"Name: {resume.name}"]
    if resume.summary:
        lines.append(f"\nSummary:\n{resume.summary}")
    if resume.skills:
        lines.append(f"\nSkills: {', '.join(resume.skills)}")
    if resume.experience:
        lines.append("\nExperience:")
        for e in resume.experience:
            end = e.end or "present"
            lines.append(f"- {e.title} at {e.company} ({e.start} to {end})")
            for b in e.bullets:
                lines.append(f"  * {b}")
    if resume.education:
        lines.append("\nEducation:")
        for ed in resume.education:
            deg = f", {ed.degree}" if ed.degree else ""
            lines.append(f"- {ed.school}{deg}")
    return "\n".join(lines)


def docx_plain_text(path: Path) -> str:
    """Extract body + table cell text from a DOCX for post-tailor scoring."""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
    return "\n".join(parts)
